# -*- coding: utf-8 -*-
"""Đọc file Excel thành lưới giá trị (JSON) cho các công cụ học / làm sạch chạy bằng Node.
   Dùng: python doc_xlsx.py <file.xlsx>   → in JSON {"sheets": [{"name", "state", "grid": [[...], ...]}]}"""
import sys, json, datetime
from openpyxl import load_workbook

wb = load_workbook(sys.argv[1], data_only=True, read_only=True)
out = []
for ws in wb.worksheets:
    grid = []
    for row in ws.iter_rows(values_only=True):
        r = []
        for v in row:
            if isinstance(v, (datetime.date, datetime.datetime, datetime.time)): v = str(v)
            elif isinstance(v, str): v = v.strip() or None
            r.append(v)
        while r and r[-1] is None: r.pop()
        grid.append(r)
    while grid and not grid[-1]: grid.pop()
    out.append({"name": ws.title, "state": ws.sheet_state, "grid": grid})
sys.stdout.write(json.dumps({"sheets": out}, ensure_ascii=False))
