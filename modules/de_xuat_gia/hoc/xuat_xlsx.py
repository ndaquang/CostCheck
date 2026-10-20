# -*- coding: utf-8 -*-
"""Ghi báo cáo Excel từ JSON do hoc.js / lam_sach.js tạo.
   JSON: {"sheets": [{"name", "title", "note", "header": [...], "widths": [...], "rows": [[...]], "pct": [chỉ số cột %], "money": [chỉ số cột tiền]}]}
   Dùng: python xuat_xlsx.py <bao_cao.json> <file.xlsx>"""
import sys, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

d = json.load(open(sys.argv[1], encoding='utf-8'))
thin = Side(style='thin', color='FFBFBFBF'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill('solid', fgColor='FF1F4E79'); HF = Font(bold=True, color='FFFFFFFF')
wb = Workbook(); wb.remove(wb.active)
for sh in d['sheets']:
    ws = wb.create_sheet(sh['name'][:31])
    r = 1
    if sh.get('title'):
        ws.cell(r, 1, sh['title']).font = Font(bold=True, size=13, color='FF1F4E79'); r += 1
    if sh.get('note'):
        for line in sh['note'].split('\n'):
            ws.cell(r, 1, line).font = Font(italic=True, color='FF595959'); r += 1
    if sh.get('title') or sh.get('note'): r += 1
    hr = r
    for c, t in enumerate(sh.get('header', []), 1):
        x = ws.cell(r, c, t); x.font = HF; x.fill = HEAD; x.border = B
        x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    pct, money = set(sh.get('pct', [])), set(sh.get('money', []))
    for row in sh.get('rows', []):
        r += 1
        for c, v in enumerate(row, 1):
            x = ws.cell(r, c, v); x.border = B; x.alignment = Alignment(vertical='top', wrap_text=isinstance(v, str) and len(v) > 40)
            if c - 1 in pct and isinstance(v, (int, float)): x.number_format = '0.0%'
            elif c - 1 in money and isinstance(v, (int, float)): x.number_format = '#,##0'
    for c, w in enumerate(sh.get('widths', []), 1): ws.column_dimensions[get_column_letter(c)].width = w
    if sh.get('header'): ws.freeze_panes = ws.cell(hr + 1, 1); ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(sh['header']))}{max(r, hr)}"
wb.save(sys.argv[2])
