# -*- coding: utf-8 -*-
"""Báo cáo kết quả điền giá (file Excel RIÊNG, không nằm trong BOQ) từ ket_qua.json do chay_file.js tạo.
   Dùng: python bao_cao_xlsx.py <ket_qua.json> <KET_QUA.xlsx> [ghi chú CSDL]"""
import sys, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.formatting.rule import FormulaRule

src, out = sys.argv[1], sys.argv[2]
ghi_chu_csdl = sys.argv[3] if len(sys.argv) > 3 else ''
d = json.load(open(src, encoding='utf-8'))
rows, cau_truc = d['rows'], d['cauTruc']
PICK = {'p25': 'Tốt nhất — cạnh tranh (P25)', 'p10': 'Thấp nhất hợp lý (P10)', 'med': 'Trung vị (P50)', 'latest': 'Năm gần nhất'}
FIT = {'cao': 'Cao', 'tb': 'Trung bình', 'thap': 'Thấp', 'none': 'Không khớp'}
EV = {'ok': 'Hợp lý', 'xem': 'Cần xem', 'bt': 'Bất thường', 'notk': 'Không có tham khảo', 'nogia': '', 'skip': ''}

thin = Side(style='thin', color='FFBFBFBF'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill('solid', fgColor='FF1F4E79'); HF = Font(bold=True, color='FFFFFFFF'); REF = PatternFill('solid', fgColor='FFF2F2F2')
COL = d.get('mode', 'col') == 'col'   # giá đề xuất ghi vào cột riêng (mặc định) hay vào ô Đơn giá
L_FILL, L_KEEP, L_NONE = ('Có giá đề xuất', 'Chỉ có giá QS', 'Để trống — QS định giá') if COL else ('Đã điền', 'Giữ giá có sẵn', 'Để trống — QS định giá')
ST = {L_FILL: ('FFE2F0D9', 'FF1F6B3A'), L_KEEP: ('FFDDEBF7', 'FF1F4E79'), L_NONE: ('FFFCE4D6', 'FF9C5700')}

def header(ws, r, titles, widths):
    for c, t in enumerate(titles, 1):
        x = ws.cell(r, c, t); x.font = HF; x.fill = HEAD; x.border = B
        x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for c, w in enumerate(widths, 1): ws.column_dimensions[ws.cell(1, c).column_letter].width = w

status = lambda r: L_FILL if r['dien'] else (L_KEEP if r.get('gia') or r.get('gia_co_san') else L_NONE)
wb = Workbook()
ws = wb.active; ws.title = 'Tóm tắt'
ws['A1'] = f"KẾT QUẢ ĐIỀN GIÁ BOQ — {d['file']}"; ws['A1'].font = Font(bold=True, size=14, color='FF1F4E79')
if ghi_chu_csdl: ws['A2'] = ghi_chu_csdl; ws['A2'].font = Font(italic=True, color='FFC00000')
n_fill = sum(1 for r in rows if r['dien']); n_keep = sum(1 for r in rows if status(r) == L_KEEP)
amt = lambda f: sum(r['gia'] * r['kl'] for r in rows if f(r) and r.get('gia') and r.get('kl'))
kpi = [
    ('File BOQ đã điền giá', d.get('out') or '(không có ô nào được điền)'),
    ('Ghi giá vào', 'Cột riêng "ĐƠN GIÁ ĐỀ XUẤT (CSDL)" tô màu, sau cột cuối của bảng — cột Đơn giá / Thành tiền giữ nguyên' if COL else 'Ô Đơn giá đang trống'),
    ('Cách chọn giá', PICK.get(d.get('pick'), d.get('pick'))),
    ('Đơn giá bổ sung đã nạp', ', '.join(d.get('extra') or []) or 'không có'),
    ('Dòng hạng mục', len(rows)),
    (L_FILL, n_fill), (L_KEEP, n_keep), (L_NONE, len(rows) - n_fill - n_keep),
    *([('Tổng giá trị theo giá đề xuất (VNĐ)', round(amt(lambda r: r['dien'])))] if COL else
      [('Tổng giá trị sau khi điền (VNĐ)', round(amt(lambda r: True))), ('   trong đó phần mới điền (VNĐ)', round(amt(lambda r: r['dien'])))]),
    ('Kiểm tra định dạng BOQ', d['fmt']),
    ('Thời gian đọc / khớp / điền (giây)', f"{d['tRead']/1000:.1f} / {d['tMatch']/1000:.1f} / {d['tExport']/1000:.1f}"),
]
header(ws, 4, ['Chỉ tiêu', 'Kết quả'], [40, 90])
for i, (k, v) in enumerate(kpi, 5):
    ws.cell(i, 1, k).border = B; c = ws.cell(i, 2, v); c.border = B; c.alignment = Alignment(wrap_text=True, horizontal='left')
    if isinstance(v, (int, float)): c.number_format = '#,##0'
    if k in ST: c.fill = PatternFill('solid', fgColor=ST[k][0]); c.font = Font(bold=True, color=ST[k][1])

ws = wb.create_sheet('Đọc file')
header(ws, 1, ['Sheet', 'Dùng', 'Dòng tiêu đề', 'Mã', 'Mô tả', 'Đơn vị', 'Khối lượng', 'Đơn giá', 'Thành tiền', 'Baseline'], [18, 8, 12, 8, 8, 8, 11, 14, 11, 10])
for i, s in enumerate(cau_truc, 2):
    c = s['cot'] if s['dung'] else {}
    for j, v in enumerate([s['sheet'], 'có' if s['dung'] else 'không', s['dong_tieu_de'] if s['dung'] else None] + [c.get(k) for k in ['code', 'desc', 'unit', 'qty', 'rate', 'amount', 'base']], 1):
        ws.cell(i, j, v).border = B

ws = wb.create_sheet('Chi tiết')
H = ['Sheet', 'Dòng', 'Mã', 'Mô tả', 'Tiêu đề gần nhất', 'ĐV', 'KL', 'Nhóm tham khảo CSDL', 'Độ phù hợp', 'Số BG',
     'P10', 'P25', 'P50', 'P75', 'P90', 'Giá đề xuất' if COL else 'Đơn giá', 'Thành tiền', 'Trạng thái', 'Lý do', 'Giá có sẵn: đánh giá', 'Lưu ý dữ liệu', 'Giá có sẵn trong BOQ']
header(ws, 1, H, [12, 6, 6, 44, 28, 7, 10, 34, 11, 7, 11, 11, 11, 11, 11, 13, 15, 20, 34, 16, 30, 14])
for i, r in enumerate(rows, 2):
    ref = r.get('ref') or {}; nhom = r.get('nhom') or ''
    vals = [r['sheet'], r['dong'], r['ma'], r['mo_ta'], r['near'], r['dv'], r['kl'], nhom, FIT.get(r['phu_hop'], ''), r.get('n'),
            *[round(ref[k]) if ref else None for k in ['p10', 'p25', 'p50', 'p75', 'p90']],
            r.get('gia'), (r['gia'] * r['kl']) if r.get('gia') and r.get('kl') else None, status(r), r['ly_do'],
            EV.get(r.get('danh_gia_co_san') or '', ''), r.get('luu_y'), r.get('gia_co_san')]
    for j, v in enumerate(vals, 1):
        c = ws.cell(i, j, v); c.border = B
        if j in (4, 5, 8, 19, 21): c.alignment = Alignment(wrap_text=True, vertical='top')
        if j == 7: c.number_format = '#,##0.##'
        if 11 <= j <= 17 or j == 22: c.number_format = '#,##0'
        if 11 <= j <= 15: c.fill = REF
        if j == 16: c.font = Font(bold=True)
n = len(rows) + 1
for lbl, (bg, fg) in ST.items():
    ws.conditional_formatting.add(f'R2:R{n}', FormulaRule(formula=[f'$R2="{lbl}"'], fill=PatternFill('solid', bgColor=bg), font=Font(color=fg, bold=True)))
ws.freeze_panes = 'E2'; ws.auto_filter.ref = f'A1:V{n}'
wb.save(out)
print('ok', out)
