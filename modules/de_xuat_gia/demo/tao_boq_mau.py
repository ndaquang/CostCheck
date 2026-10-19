# -*- coding: utf-8 -*-
"""BOQ dự án GIẢ LẬP để chạy thử module (không phải dữ liệu thật).
   Dùng: python tao_boq_mau.py <thư mục ra>
   BOQ_ChayThu.xlsx     BOQ chưa có giá (định dạng như BOQ thật: tiêu đề song ngữ 2 dòng, gộp ô, màu, viền,
                        công thức thành tiền, sheet Tổng hợp tham chiếu sang các sheet chi tiết, vùng in).
   BOQ_ChayThu_QS.xlsx  cùng BOQ, QS đã nhập đơn giá.
   Baseline_CDT.xlsx    giá baseline của chủ đầu tư cho một phần hạng mục."""
import sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

OUT = sys.argv[1] if len(sys.argv) > 1 else '.'
os.makedirs(OUT, exist_ok=True)
thin = Side(style='thin', color='FF808080'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill('solid', fgColor='FF1F4E79'); SUB = PatternFill('solid', fgColor='FFDDEBF7'); GRP = PatternFill('solid', fgColor='FFFCE4D6')
WHITE = Font(bold=True, color='FFFFFFFF', name='Arial', size=10)

# (mã, mô tả, đơn vị, khối lượng, giá QS, baseline CĐT) — giá QS None = QS chưa nhập; mô tả viết khác cách ghi trong CSDL
KET_CAU = [
    ('A', 'CÔNG TÁC ĐẤT VÀ MÓNG', None, None, None, None),
    ('1', 'Đào đất hố móng bằng máy đào, đất cấp II', 'm³', 1250, 51000, 45000),
    ('2', 'Đắp đất hố móng, đầm chặt K=0,90', 'm³', 680, 60000, 58000),
    ('3', 'Vận chuyển đất thừa đổ đi, cự ly 10km', 'm³', 570, 145000, 90000),          # cao bất thường
    ('4', 'Ép cọc BTCT tiết diện 300x300', 'md', 4200, 182000, 175000),
    ('5', 'Bê tông lót móng đá 4x6, mác 150, dày 100', 'm³', 48.5, 1150000, 1100000),
    ('6', 'Bê tông đài móng, giằng móng mác 300 đá 1x2', 'm³', 412, 1690000, 1650000),
    ('7', 'Cốt thép móng đường kính D≤10', 'tấn', 8.6, 22300000, 20500000),
    ('8', 'Cốt thép móng đường kính D≤18', 'tấn', 31.2, 20900000, 20000000),
    ('9', 'Ván khuôn móng, giằng móng', 'm²', 1380, 120000, 175000),                 # thấp bất thường
    ('B', 'KẾT CẤU THÂN', None, None, None, None),
    ('10', 'Bê tông cột mác 300 đá 1x2', 'm³', 96, 1880000, 1800000),
    ('11', 'Bê tông dầm, sàn mác 300 đá 1x2', 'm³', 865, 1760000, 1700000),
    ('12', 'Cốt thép cột D>18', 'tấn', 18.4, 21800000, 20000000),
    ('13', 'Cốt thép dầm sàn D≤10', 'tấn', 42.5, 21500000, 21000000),
    ('14', 'Ván khuôn cột', 'm²', 760, 255000, 220000),
    ('15', 'Ván khuôn dầm sàn', 'm²', 5600, 210000, 205000),
    ('16', 'Nền bê tông cốt thép dày 200, M300', 'm²', 3200, 560000, 480000),
    ('17', 'Kết cấu thép khung nhà xưởng (gia công + lắp dựng)', 'tấn', 125, 33500000, 31000000),
    ('18', 'Sơn chống cháy kết cấu thép R60', 'm²', 4100, None, 270000),               # QS chưa nhập
    ('19', 'Lợp mái tôn mạ màu dày 0,45mm', 'm²', 5400, 228000, 230000),
]
HOAN_THIEN = [
    ('C', 'XÂY TÔ', None, None, None, None),
    ('20', 'Xây tường gạch đặc dày 220, vữa xi măng M75', 'm³', 145, 1510000, 1450000),
    ('21', 'Xây tường gạch ống dày 100', 'm²', 980, 330000, 250000),                   # cao (cần xem/bất thường)
    ('22', 'Trát tường trong, vữa M75 dày 15mm', 'm²', 4300, 99000, 95000),
    ('23', 'Trát tường ngoài, vữa M75 dày 20mm', 'm²', 2100, 128000, 120000),
    ('D', 'HOÀN THIỆN', None, None, None, None),
    ('24', 'Bả matit tường trong nhà', 'm²', 4300, 36000, 35000),
    ('25', 'Sơn nước tường trong nhà, 1 lớp lót + 2 lớp phủ', 'm²', 4300, 69000, 65000),
    ('26', 'Sơn nước tường ngoài nhà, 1 lớp lót + 2 lớp phủ', 'm²', 2100, 98000, 90000),
    ('27', 'Trần thạch cao khung chìm, tấm dày 9mm', 'm²', 1850, 262000, 255000),
    ('28', 'Trần thạch cao chống ẩm khu WC', 'm²', 240, 305000, None),
    ('29', 'Lát nền gạch granite 800x800', 'm²', 1620, 470000, 480000),
    ('30', 'Ốp tường WC gạch ceramic 300x600', 'm²', 520, 340000, None),
    ('31', 'Chống thấm sàn WC bằng màng khò nóng', 'm²', 260, 232000, 220000),
    ('32', 'Cửa đi nhôm kính cường lực 8mm', 'm²', 145, 2700000, 2500000),
    ('33', 'Vách ngăn WC compact HPL 12mm', 'm²', 96, 1420000, None),
    ('34', 'Lan can cầu thang inox 304', 'md', 85, 1180000, 1100000),
    ('E', 'HẠNG MỤC KHÁC (MEP — ngoài CSDL)', None, None, None, None),
    ('35', 'Máy bơm nước sinh hoạt Q=10m³/h', 'bộ', 2, 38500000, None),
    ('36', 'Tủ điện tổng MSB 800A', 'tủ', 1, 285000000, None),
]

def sheet(wb, title, ten, rows, with_price):
    ws = wb.create_sheet(title)
    ws.merge_cells('A1:H1'); ws['A1'] = ten; ws['A1'].font = Font(bold=True, size=14, name='Arial', color='FF1F4E79')
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A2:H2'); ws['A2'] = 'Dự án: NHÀ XƯỞNG GIẢ LẬP — Gói thầu: Xây dựng phần thô và hoàn thiện'
    ws['A2'].font = Font(italic=True, name='Arial', size=10); ws['A2'].alignment = Alignment(horizontal='center')
    h1 = ['STT', 'Mô tả công việc', 'Đơn vị', 'Khối lượng', 'Đơn giá (VNĐ)', 'Thành tiền (VNĐ)', 'Ghi chú', 'Ref']
    h2 = ['Item', 'Description', 'Unit', 'Quantity', 'Unit rate', 'Amount', 'Remark', '']
    for c, v in enumerate(h1, 1):
        x = ws.cell(4, c, v); x.font = WHITE; x.fill = HEAD; x.border = B; x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for c, v in enumerate(h2, 1):
        x = ws.cell(5, c, v); x.font = Font(italic=True, name='Arial', size=9); x.fill = SUB; x.border = B; x.alignment = Alignment(horizontal='center')
    r = 6; first = r
    for code, desc, unit, qty, qs, _ in rows:
        ws.cell(r, 1, code); ws.cell(r, 2, desc)
        if unit:
            ws.cell(r, 3, unit); ws.cell(r, 4, qty)
            if with_price and qs is not None: ws.cell(r, 5, qs)
            ws.cell(r, 6, f'=IF(E{r}="","",D{r}*E{r})')
            for c in (4,): ws.cell(r, c).number_format = '#,##0.00'
            for c in (5, 6): ws.cell(r, c).number_format = '#,##0'
            for c in range(1, 9): ws.cell(r, c).font = Font(name='Arial', size=10)
        else:
            for c in range(1, 9): ws.cell(r, c).fill = GRP; ws.cell(r, c).font = Font(bold=True, name='Arial', size=10)
        for c in range(1, 9): ws.cell(r, c).border = B
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')
        r += 1
    ws.cell(r, 2, 'TỔNG CỘNG'); ws.cell(r, 6, f'=SUM(F{first}:F{r-1})')
    for c in range(1, 9): ws.cell(r, c).font = Font(bold=True, name='Arial', size=10); ws.cell(r, c).border = B; ws.cell(r, c).fill = SUB
    ws.cell(r, 6).number_format = '#,##0'
    for col, w in zip('ABCDEFGH', [7, 52, 8, 12, 16, 18, 18, 6]): ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A6'; ws.print_area = f'A1:H{r}'; ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1
    return r

def build(path, with_price):
    wb = Workbook(); th = wb.active; th.title = 'Tong hop'
    t1 = sheet(wb, '1. Ket cau', 'BẢNG KHỐI LƯỢNG — PHẦN KẾT CẤU', KET_CAU, with_price)
    t2 = sheet(wb, '2. Hoan thien', 'BẢNG KHỐI LƯỢNG — PHẦN HOÀN THIỆN', HOAN_THIEN, with_price)
    th['A1'] = 'BẢNG TỔNG HỢP GIÁ TRỊ GÓI THẦU'; th['A1'].font = Font(bold=True, size=14, name='Arial', color='FF1F4E79')
    for c, v in enumerate(['STT', 'Hạng mục', 'Giá trị (VNĐ)'], 1):
        x = th.cell(3, c, v); x.font = WHITE; x.fill = HEAD; x.border = B
    th.append([1, 'Phần kết cấu', f"='1. Ket cau'!F{t1}"]); th.append([2, 'Phần hoàn thiện', f"='2. Hoan thien'!F{t2}"])
    th.append([None, 'TỔNG CỘNG (chưa VAT)', '=SUM(C4:C5)'])
    for r in range(4, 7):
        for c in range(1, 4): th.cell(r, c).border = B
        th.cell(r, 3).number_format = '#,##0'
    th['B6'].font = Font(bold=True); th.column_dimensions['B'].width = 36; th.column_dimensions['C'].width = 22
    wb.save(path)

build(os.path.join(OUT, 'BOQ_ChayThu.xlsx'), False)
build(os.path.join(OUT, 'BOQ_ChayThu_QS.xlsx'), True)
bw = Workbook(); bs = bw.active; bs.title = 'Baseline'
bs.append(['STT', 'Mô tả công việc', 'Đơn vị', 'Đơn giá baseline (VNĐ)'])
for rows in (KET_CAU, HOAN_THIEN):
    for code, desc, unit, qty, qs, base in rows:
        if unit and base: bs.append([code, desc, unit, base])
bw.save(os.path.join(OUT, 'Baseline_CDT.xlsx'))
print('ok')
